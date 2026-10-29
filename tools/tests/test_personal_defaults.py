import json,subprocess,sys,tempfile,unittest
from pathlib import Path
from datetime import date,timedelta
import openpyxl
from openpyxl.worksheet.table import Table

class PersonalDefaultsTest(unittest.TestCase):
 def test_sdf_uses_national_regional_and_local_holidays_without_changing_rotating_groups(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);book=openpyxl.Workbook();s=book.active;s.append(['DNE','DSO']);s.append(['99999','SDF']);s.add_table(Table(displayName='Cuadro_2027',ref='A1:B2'))
   cal=book.create_sheet('Calendario');cal.append(['Fecha','GR1','GR2','GR3','GR4','GR1N','GR2N','GR3N','GR4N'])
   for i in range(365):cal.append([date(2027,1,1)+timedelta(days=i),'T','T','T','A','T','T','T','A'])
   cal.add_table(Table(displayName='Calendario_Diario_2027',ref='A1:I366'));book.save(root/'book.xlsx')
   data={'schema':'scmm-personal-simulator-v1','entries':[{'dne':'99999','years':{}}]};(root/'profiles.json').write_text(json.dumps(data))
   script=Path(__file__).resolve().parents[1]/'prepare_personal_defaults.py'
   subprocess.run([sys.executable,str(script),str(root/'profiles.json'),str(root/'book.xlsx'),str(root/'out.json')],check=True,capture_output=True)
   result=json.loads((root/'out.json').read_text());calendar=result['calendars']['2027'];rests=calendar['groups']['SDF']
   self.assertEqual(len(rests),115);self.assertEqual(len(set(rests)),115);self.assertEqual(len(calendar['holidays']),14)
   for d in ['2027-01-01','2027-03-19','2027-03-25','2027-03-26','2027-05-15','2027-08-16','2027-11-09','2027-12-06','2027-12-08','2027-12-25','2027-01-02','2027-01-03']:self.assertIn(d,rests)
   self.assertNotIn('2027-05-17',rests) # No invented Monday transfer of San Isidro.
   self.assertNotIn('2027-11-10',rests)
   self.assertEqual(calendar['groups']['1'],[]) # Holidays aren't forced rests in rotating groups.
   self.assertEqual(calendar['groups']['4N'],[])
   self.assertEqual(set(result['calendars']),{'2027'}) # No calendar projected into another year.
if __name__=='__main__':unittest.main()
