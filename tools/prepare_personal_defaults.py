#!/usr/bin/env python3
"""Join private MTE vacation profiles and the validated annual workbook.

Only the explicitly selected output is written; no nominal workbook fields are
copied. Keep input/output outside the repository. Preview packaging is separate.
"""
import argparse,calendar,hashlib,json,re
from pathlib import Path
from datetime import date,timedelta,datetime
from collections import Counter
import openpyxl
from openpyxl.utils import range_boundaries
p=argparse.ArgumentParser();p.add_argument('profiles',type=Path);p.add_argument('workbook',type=Path);p.add_argument('output',type=Path);a=p.parse_args()
repo=Path(__file__).resolve().parents[1]
if a.output.resolve().is_relative_to(repo):p.error('Write assignment data outside the repository')
x=json.loads(a.profiles.read_text());assert x['schema']=='scmm-personal-simulator-v1'
w=openpyxl.load_workbook(a.workbook,data_only=True)
def table(name):
 for s in w:
  if name in s.tables:
   left,top,right,bottom=range_boundaries(s.tables[name].ref)
   return [[s.cell(r,c).value for c in range(left,right+1)] for r in range(top,bottom+1)]
 raise ValueError('Missing table '+name)
rows=table('Cuadro_2027');headers=rows.pop(0);assert 'DNE' in headers and 'DSO' in headers
assigned={}
for cells in rows:
 row=dict(zip(headers,cells));id=str(row['DNE']).strip();g=row['DSO']
 assert re.fullmatch(r'\d{1,10}',id) and id not in assigned
 assert g is None or re.fullmatch(r'GR[1-4]N?|SDF',g)
 assigned[id]=int(g[-1]) if g and re.fullmatch(r'GR[1-4]',g) else g.removeprefix('GR') if g else None
for e in x['entries']:
 e['restGroup']=assigned.get(e['dne'])
 for plan in e['years'].values():plan['sourceReference']='Vacaciones y fiestas agrupadas MTE 2027; asignación de descanso: cuadro anual 2027.'
rows=table('Calendario_Diario_2027');headers=rows.pop(0);assert headers==['Fecha','GR1','GR2','GR3','GR4','GR1N','GR2N','GR3N','GR4N'];assert len(rows)==365
rest={g.removeprefix('GR'):[] for g in headers[1:]}
for i,row in enumerate(rows):
 d=row[0].date() if isinstance(row[0],datetime) else date.fromisoformat(str(row[0])[:10])
 assert d==date(2027,1,1)+timedelta(days=i)
 for g,value in zip(headers[1:],row[1:]):
  assert value in ('ABCD' if g.startswith('GR4') else 'TD')
  if value=='D':rest[g.removeprefix('GR')].append(d.isoformat())
holiday_ref=json.loads((Path(__file__).parent/'data/madrid-festivos-2027.json').read_text())
assert holiday_ref['year']==2027 and holiday_ref['verified'] is True
holidays=set(holiday_ref['dates']);assert len(holidays)==14
for d in holidays:assert date.fromisoformat(d).year==2027
rest['SDF']=[d.isoformat() for i in range(365) if (d:=date(2027,1,1)+timedelta(days=i)).weekday()>=5 or d.isoformat() in holidays]
x['calendars']={'2027':{'groups':rest,'holidays':sorted(holidays),'sourceReference':'Calendario_Diario_2027 del libro de cuadros. SDF: sábados, domingos y fiestas laborales del municipio de Madrid (regla confirmada por el usuario). '+holiday_ref['sourceReference']+' '+holiday_ref['sourceUrl'],'verified':True}}
x['version']='2027-vacaciones-escalafon-descansos-sdf-20261028';x['sourceReference']='Escalafón conciliado, vacaciones/PFA y cuadro anual MTE 2027; sin nombres ni DNI.'
a.output.parent.mkdir(parents=True,exist_ok=True);a.output.write_text(json.dumps(x,ensure_ascii=False,separators=(',',':')))
print(json.dumps({'profiles':len(x['entries']),'assigned':sum(e['restGroup'] is not None for e in x['entries']),'missing':sum(e['restGroup'] is None for e in x['entries']),'sourceSha256':hashlib.sha256(a.workbook.read_bytes()).hexdigest(),'groups':dict(Counter(str(e['restGroup']) for e in x['entries']))}))
